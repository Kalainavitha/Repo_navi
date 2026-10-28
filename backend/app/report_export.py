"""Download formats for the weekly report: Word (.docx), PDF and CSV, built from the same report dictionary."""
import csv
import html
import io

from .report import _day, _hours, _nice

EXPORT_FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pdf": ("application/pdf", "pdf"),
    "csv": ("text/csv; charset=utf-8", "csv"),
}


def _sections(report: dict) -> list[tuple[str, list[str]]]:
    """The report as (heading, lines) pairs, shared by the Word and PDF exports."""
    who = lambda t: f" ({t['assignee_name']})" if t.get("assignee_name") else ""
    due = lambda t: f", due {_nice(_day(t.get('end_date')))}" if t.get("end_date") else ""
    blocked_ids = {b["id"] for b in report["blocked"]}
    attention = report["blocked"] + [t for t in report["overdue"] if t["id"] not in blocked_ids]
    progress = report["progress"]
    overview = [
        f"Overall: {report['status']} ({'; '.join(report['status_reasons'])})",
        f"Progress: {progress['done']}/{progress['total']} tasks done ({progress['percent']}%), "
        f"{_hours(progress['hours_done'])} of {_hours(progress['hours_total'])}, {_hours(progress['hours_remaining'])} remaining",
    ]
    if report.get("forecast"):
        overview.append(f"Forecast: {report['forecast']}")
    return [
        ("Overview", overview),
        ("Completed this week", [f"{t['title']}{who(t)}" for t in report["completed_this_week"]] or ["Nothing completed this week."]),
        ("In progress", [f"{t['title']}{who(t)}{due(t)}" for t in report["in_progress"]] or ["Nothing in progress."]),
        ("Blocked or overdue", [f"{'HALTED' if t['status'] == 'Blocked' else 'OVERDUE'}: {t['title']}{who(t)}{due(t)}" for t in attention] or ["None."]),
        (f"Next week ({report['next_week_label']})", [
            f"{p['name']}: {_hours(p['hours'])} - " + "; ".join(p["tasks"][:3]) + (" ..." if len(p["tasks"]) > 3 else "")
            for p in report["next_week"]
        ] or ["No work is scheduled."]),
        ("Top risks", [f"[{r['severity']}] {r['message']}" for r in report["risks"]] or ["No significant risks."]),
    ]


def to_docx(report: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(f"Weekly status: {report['project']}", level=0)
    document.add_paragraph(f"Week: {report['week_label']}")
    for heading, lines in _sections(report):
        document.add_heading(heading, level=2)
        for line in lines:
            paragraph = document.add_paragraph(line, style=None if heading == "Overview" else "List Bullet")
            if line.startswith(("HALTED", "OVERDUE")):
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(report: dict) -> bytes:
    import pymupdf

    tone = {"On track": "#245C43", "At risk": "#B45309", "Late": "#B42318"}.get(report["status"], "#245C43")
    parts = [
        f"<h1 style='color:#123D2B'>Weekly status: {html.escape(report['project'])}</h1>",
        f"<p style='color:#555555'>Week: {html.escape(report['week_label'])}</p>",
        f"<p style='color:{tone}'><b>{html.escape(report['status'])}</b></p>",
    ]
    for heading, lines in _sections(report):
        parts.append(f"<h3 style='color:#245C43'>{html.escape(heading)}</h3>")
        parts.append("<ul>" + "".join(f"<li>{html.escape(line)}</li>" for line in lines) + "</ul>")
    story = pymupdf.Story(html="<body style='font-family: sans-serif; font-size: 11pt'>" + "".join(parts) + "</body>")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page_box = pymupdf.paper_rect("a4")
    content_box = page_box + (50, 50, -50, -50)
    more = 1
    while more:
        device = writer.begin_page(page_box)
        more, _filled = story.place(content_box)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()


def to_csv(report: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Report", "Week", "Section", "Item", "Owner", "Status", "Due", "Hours"])
    base = [report["project"], report["week_label"]]
    progress = report["progress"]
    writer.writerow(base + ["Overview", "Overall status", "", report["status"], report.get("deadline") or "", ""])
    writer.writerow(base + ["Overview", "Progress", "", f"{progress['percent']}% ({progress['done']}/{progress['total']} tasks)", "", progress["hours_done"]])
    writer.writerow(base + ["Overview", "Hours remaining", "", "", "", progress["hours_remaining"]])
    if report.get("forecast"):
        writer.writerow(base + ["Overview", report["forecast"], "", "", "", ""])
    blocked_ids = {b["id"] for b in report["blocked"]}
    for label, key in (("Completed this week", "completed_this_week"), ("In progress", "in_progress"), ("Blocked", "blocked"), ("Overdue", "overdue")):
        for task in report[key]:
            if key == "overdue" and task["id"] in blocked_ids:
                continue
            writer.writerow(base + [label, task["title"], task.get("assignee_name") or "", task["status"], task.get("end_date") or "", ""])
    for person in report["next_week"]:
        writer.writerow(base + [f"Next week ({report['next_week_label']})", "; ".join(person["tasks"]), person["name"], "Planned", "", person["hours"]])
    for risk in report["risks"]:
        writer.writerow(base + ["Top risks", risk["message"], "", risk["severity"], "", ""])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM so Excel shows accents correctly


def export_report(report: dict, fmt: str) -> tuple[bytes, str, str]:
    """Returns (file bytes, media type, file extension)."""
    media, extension = EXPORT_FORMATS[fmt]
    builders = {"docx": to_docx, "pdf": to_pdf, "csv": to_csv}
    return builders[fmt](report), media, extension
